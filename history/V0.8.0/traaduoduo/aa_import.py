"""Read-only AA import. Formula caches are evidence, never executable code."""
import hashlib
import uuid
import posixpath
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET

HEADERS=['檔別','動支單號','工程項目','工程項目中文說明','計算來源項目','計算來源項目中文說明','序號','本局材料編號','工程數量單位','數量','單價','金額','計算用比例','算式說明','規格說明','備註']
TABLES={'budget':'預算表','tender':'招標表','contract':'契約表','changes':'變更表'}
NS={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}

def format_decimal(value):
    return '' if value=='' or value is None else f"{Decimal(str(value)).quantize(Decimal('.01'),rounding=ROUND_HALF_UP):,.2f}"

def sheets(path):
    from openpyxl import load_workbook
    with Path(path).open('rb') as stream:
        wb=load_workbook(stream,read_only=True,keep_links=False)
        try: return wb.sheetnames
        finally: wb.close()

def inspect_aa(path,sheet=None,header_row=1):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    path=Path(path).resolve()
    if path.suffix.lower()!='.xlsx': raise ValueError('請選擇 .xlsx 格式的 AA 檔案')
    digest=hashlib.sha256(path.read_bytes()).hexdigest()
    with ZipFile(path) as z:
        if sum(i.file_size for i in z.infolist())>150_000_000: raise ValueError('活頁簿解壓後過大，請先拆分資料')
        external=sum(n.startswith('xl/externalLinks/externalLink') and n.endswith('.xml') for n in z.namelist())
        cache={}; raw_formulas={}
        book=ET.fromstring(z.read('xl/workbook.xml'))
        sheet_info=list(book.find('s:sheets',NS))
        chosen=next((s for s in sheet_info if s.get('name')==(sheet or sheet_info[0].get('name'))),None)
        if chosen is None: raise ValueError('找不到工作表')
        sheet=chosen.get('name')
        rid=chosen.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        rels=ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
        target=next(r.get('Target') for r in rels if r.get('Id')==rid)
        xml_path=target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/'+target)
        raw=ET.fromstring(z.read(xml_path))
        for c in raw.findall('.//s:c',NS):
            v=c.find('s:v',NS); f=c.find('s:f',NS)
            if v is not None and c.get('t') not in ('s','str','inlineStr','b','e'): cache[c.get('r')]=v.text or ''
            if f is not None: raw_formulas[c.get('r')]=dict(text=f.text or '',attributes=dict(f.attrib))
    with path.open('rb') as stream:
        wb=load_workbook(stream,read_only=True,data_only=False,keep_links=False)
        try:
            ws=wb[sheet]
            if ws.max_row>100_000: raise ValueError('單次匯入上限 100,000 列，請先拆分')
            headers=[str(c.value or '').strip() for c in next(ws.iter_rows(min_row=header_row,max_row=header_row,max_col=16))]
            if headers!=HEADERS: raise ValueError('前 16 欄表頭或順序不符 AA 格式；請確認工作表與表頭列')
            rows=[]; warnings=[]; errors=[]
            for source_row,cells in enumerate(ws.iter_rows(min_row=header_row+1),start=header_row+1):
                if not any(c.value is not None for c in cells): continue
                values=[]; formulas={}; raw_values={}; extra={}
                for i,c in enumerate(cells):
                    coordinate=getattr(c,'coordinate',get_column_letter(i+1)+str(source_row))
                    value=c.value
                    if c.data_type=='f':
                        formulas[coordinate]=dict(formula=str(value),raw=raw_formulas.get(coordinate,{}),cache=cache.get(coordinate,''))
                        value=cache.get(coordinate,'')
                        if value=='': warnings.append(f'{coordinate} 公式無快取值，未補 0')
                    elif coordinate in cache: value=cache[coordinate]
                    if value is None: value=''
                    if i in (0,2,4,7) and c.data_type=='n' and str(c.number_format).strip('0')=='':
                        if value!='': value=str(value).zfill(len(c.number_format))
                    if c.data_type=='e': warnings.append(f'{coordinate} 原檔錯誤：{value}')
                    text=str(value)
                    raw_values[coordinate]=text
                    if i<16: values.append(text)
                    elif text: extra[coordinate]=text
                values=(values+['']*16)[:16]
                if not values[2]: errors.append(f'第 {source_row} 列缺少工程項目編碼')
                for i in (9,10,11,12):
                    if values[i]:
                        try:
                            if not Decimal(values[i]).is_finite(): raise ValueError()
                        except (ValueError,ArithmeticError): errors.append(f'第 {source_row} 列 {HEADERS[i]} 不是有效數值')
                rows.append(dict(id=uuid.uuid4().hex,source_row=source_row,values=values,formulas=formulas,raw_values=raw_values,extra=extra))
        finally: wb.close()
    groups=defaultdict(dict)
    for row in rows:
        v=row['values']; scope=(v[0],v[1]); code=v[2]
        if code in groups[scope]: errors.append(f'重複工程項目：{code}')
        groups[scope][code]=row
    maximum=max((len(r['values'][2]) for r in rows),default=0)
    parent_ids=set()
    for row in rows:
        v=row['values']; code=v[2]; parents=groups[(v[0],v[1])]
        prefix=next((code[:i] for i in range(len(code)-1,0,-1) if code[:i] in parents),'')
        parent=parents[prefix]['id'] if prefix else ''
        row.update(parent_id=parent,length=len(code),analysis=len(code)==maximum)
        if parent: parent_ids.add(parent)
        elif code and len(code)>min(map(len,parents)): warnings.append(f'{code} 缺少前綴父項，保留為根節點')
    for row in rows: row['leaf']=row['id'] not in parent_ids
    if external: warnings.append(f'來源含 {external} 組外部連結定義；未更新或讀取外部檔案')
    if any(r['extra'] for r in rows): warnings.append('無標題或額外欄位已保留於來源紀錄，不會丟棄')
    if hashlib.sha256(path.read_bytes()).hexdigest()!=digest:raise ValueError('來源檔在讀取中變更，請重新預覽匯入')
    return dict(id=uuid.uuid4().hex,filename=path.name,sha256=digest,sheet=sheet,header_row=header_row,rows=rows,max_length=maximum,external_links=external,warnings=list(dict.fromkeys(warnings)),errors=errors)
