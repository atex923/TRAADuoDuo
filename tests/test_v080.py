"""V0.8 acceptance with generated files and isolated databases only."""
from copy import deepcopy
from datetime import date
from decimal import Decimal
from pathlib import Path
from threading import Event
from unittest.mock import patch
import time,unittest
from PySide6.QtCore import Qt,QDate,QEvent
from PySide6.QtGui import QInputMethodEvent
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QLineEdit,QDialog,QDialogButtonBox,QSpinBox,QTableWidgetItem,QSystemTrayIcon
from openpyxl import load_workbook
from pypdf import PdfReader
import test_ui as support
from traaduoduo.finance import synchronize,management_fee
from traaduoduo.billing import new_record,summary,effective_period,stage_summary,normalize_billing
from traaduoduo.store import ProjectStore,default_pages
from traaduoduo.report_export import Report,Column,atomic_export

def transfer(stage='first',period=1):return dict(new_record('other_payments','transfer-'+stage,period),item='transfer',stage=stage,amount='999')

class Rules(unittest.TestCase):
    def test_new_bases_and_ordered_totals(self):
        p=default_pages('公式');c=p['costs']
        for kind in ('budget','contract'):
            c[kind].update(work='1000000',work_tax='50000',award_total='1200000',labor='5000000',own_labor='100',purchased_material='200',road_stock='300',road_purchase='400',supervision='500',freight='600',pollution='700',unfinished_tax='50')
        c['extensions']=[dict(id='extra',label='其他擴充',budget='800',contract='900')]
        p=synchronize(p)
        for kind,main,other in [('budget',1050000,1500),('contract',1200000,1600)]:
            side=p['costs'][kind];self.assertEqual(management_fee(side,kind)['source'],Decimal(5001000));self.assertEqual(side['management_amount'],'150000')
            expected=Decimal(main+150000+1000+500+600+other);self.assertEqual(Decimal(side['unfinished']),expected-50);self.assertEqual(Decimal(side['total']),expected)
        self.assertEqual(p['costs']['budget']['award_total'],'1050000')
    def test_manual_management_used_once_and_missing_not_zero(self):
        p=default_pages('人工');p['costs']['budget'].update(work='100',work_tax='5',labor='0',unfinished_tax='0',management_manual=True,management_adjustment='123.456')
        q=synchronize(p);self.assertEqual(q['costs']['budget']['unfinished'],'228.456');self.assertEqual(q['costs']['budget']['total'],'228.456')
        p['costs']['budget']['work_tax']='';self.assertEqual(synchronize(p)['costs']['budget']['award_total'],'')
    def test_auto_stage_period_and_inequality_allowed(self):
        p=default_pages('階段');p['other_payments']=[transfer(period=3)];a=p['admin_payments']['allocations'][0];a.update(period=99,**{'0C12':'10','0C11':'20','0C14':'30'})
        self.assertEqual(effective_period(p,'first'),3);self.assertEqual(summary(p,'admin_payments',2)['0C12_available'],0);self.assertEqual(summary(p,'admin_payments',3)['0C12_available'],10);self.assertNotIn('提撥與分攤不一致',stage_summary(p)[0]['warnings'])
        p['other_payments'][0]['period']=5;self.assertEqual(summary(p,'admin_payments',3)['0C12_available'],0)
        p['other_payments'].append(dict(transfer(),id='duplicate'));self.assertIsNone(effective_period(p,'first'));self.assertEqual(summary(p,'admin_payments')['0C12_available'],0)
    def test_legacy_period_is_preserved_but_not_authoritative(self):
        p={'admin_payments':dict(rows=[],allocations=[dict(stage='first',period=8,**{'0C12':'1','0C11':'2','0C14':'3'})],available_basis='period')}
        q=normalize_billing(p);self.assertEqual(q['admin_payments']['allocations'][0]['historical_effective_period'],8);self.assertNotIn('allocation_revision',p['admin_payments'])

class UI(unittest.TestCase):
    setUp=support.UITests.setUp
    tearDown=support.UITests.tearDown
    wait_io=support.UITests.wait_io
    def test_idle_59_then_60_and_focus(self):
        w=self.window;w.toggle_lock('basic');f=w.fields['number'];f.setFocus();QTest.keyClicks(f,'IDLE');w.input_activity.last_input=time.monotonic()-59;w.idle_timeout();self.assertIsNone(w._auto_job);self.assertEqual(ProjectStore(self.path).load()['pages']['basic']['number'],'')
        w.input_activity.last_input-=2;position=f.cursorPosition();w.idle_timeout();self.assertTrue(w.body.isEnabled());self.wait_io();self.assertEqual(f.cursorPosition(),position);self.assertTrue(f.hasFocus());self.assertEqual(ProjectStore(self.path).load()['pages']['basic']['number'],'IDLE')
    def test_slow_background_new_typing_stays_dirty(self):
        w=self.window;w.toggle_lock('basic');f=w.fields['number'];f.setFocus();QTest.keyClicks(f,'A');entered=Event();release=Event();real=w.store.save
        def slow(*args):entered.set();release.wait(4);return real(*args)
        with patch.object(w.store,'save',side_effect=slow):
            w.start_background_save();self.assertTrue(entered.wait(1));QTest.keyClicks(f,'B');w.capture_basic();self.assertTrue(w.body.isEnabled());release.set();self.wait_io()
        self.assertEqual(ProjectStore(self.path).load()['pages']['basic']['number'],'A');self.assertEqual(w.session.pages['basic']['number'],'AB');self.assertTrue(w.session.dirty());self.assertEqual(f.text(),'AB')
        w.start_background_save();self.wait_io();self.assertFalse(w.session.dirty())
    def test_input_method_and_dialog_defer_without_committing(self):
        w=self.window;w.toggle_lock('basic');f=w.fields['number'];f.setFocus();support.APP.sendEvent(f,QInputMethodEvent('中文',[]));w.input_activity.last_input-=61;w.idle_timeout();self.assertTrue(w.input_activity.composing);self.assertIsNone(w._auto_job);support.APP.sendEvent(f,QInputMethodEvent('',[]));self.assertFalse(w.input_activity.composing)
        with patch('traaduoduo.idle_save.QApplication.activeModalWidget',return_value=object()):w.start_background_save();self.assertIsNone(w._auto_job)
    def test_auto_failure_nonblocking_and_manual_save(self):
        w=self.window;w.toggle_lock('basic');w.fields['number'].setText('draft');w.capture_basic()
        with patch.object(w.store,'save',side_effect=OSError('disk full')),patch('traaduoduo.ui.QMessageBox.warning') as warning:
            w.start_background_save();self.wait_io();self.assertFalse(warning.called)
        self.assertTrue(w.session.dirty());self.assertTrue(w.body.isEnabled());self.assertIn('失敗',w.message.text());self.assertEqual(w.revision,0);w.save();self.wait_io();self.assertFalse(w.session.dirty())
    def test_diary_search_manual_moves_date_order_and_ids(self):
        w=self.window;w.session.unlocked.add('diary');rows=[dict(id=rid,number=n,date=day,description=text,note='') for rid,n,day,text in [('a',10,'2026-10-22','材料'),('b',1,'2026-10-20','施工'),('c',2,'2026-10-22','材料備查')]];w.commit_extra('diary',rows)
        self.assertEqual(w.diary_table.columnCount(),3);self.assertEqual(w.diary_table.item(0,0).data(Qt.ItemDataRole.UserRole),'a');w.diary_table.setCurrentCell(0,1);w.move_diary(1);self.assertEqual([r['id'] for r in w.session.pages['diary']],['b','a','c'])
        w.diary_search.setText('材料');self.assertEqual(w.diary_table.rowCount(),2);before=deepcopy(w.session.pages);w.move_diary(1);self.assertEqual(w.session.pages,before);w.diary_table.item(0,2).setText('已查核');self.assertEqual(next(r for r in w.session.pages['diary'] if r['id']=='a')['note'],'已查核')
        w.diary_search.clear();w.diary_sort.setCurrentIndex(1);self.assertEqual(w.diary_table.item(0,0).data(Qt.ItemDataRole.UserRole),'b');w.diary_sort.setCurrentIndex(0);w.store.save(w.session.pages,0);self.assertEqual([r['id'] for r in w.store.load()['pages']['diary']],['b','a','c'])
    def test_header_height_and_allocation_row_height(self):
        w=self.window;w.tabs.setCurrentIndex(2);w.billing_tabs.setCurrentIndex(2);support.APP.processEvents()
        for table in [w.payment_tables['admin_payments'],w.allocation_table,w.diary_table]:
            h=table.horizontalHeader();self.assertEqual(h.height(),h.fontMetrics().height()+round(3*h.logicalDpiY()/72))
        self.assertEqual(w.allocation_table.rowHeight(0),round(18*w.logicalDpiY()/72));self.assertEqual(w.payment_summaries['admin_payments'].font().pointSize(),16)
        block=w.payment_summaries['admin_payments'].document().begin();self.assertAlmostEqual(block.blockFormat().lineHeight(),.3*w.logicalDpiY()/72,places=3)
        self.assertEqual(w.cost_fields['budget']['management_amount'].places,2);self.assertTrue(w.cost_fields['budget']['award_total'].isReadOnly())
        field=w.fields['number'];self.assertEqual(field.height(),max(field.fontMetrics().height(),field.sizeHint().height()-round(3*field.logicalDpiY()/72)))
    def test_about_content_and_non_windows_no_tray(self):
        from traaduoduo import DISPLAY_VERSION
        w=self.window;d=w.about_dialog();text=' '.join(label.text() for label in d.findChildren(__import__('PySide6.QtWidgets',fromlist=['QLabel']).QLabel));self.assertEqual(d.windowTitle(),'關於');self.assertIn(DISPLAY_VERSION,text);self.assertIn('Atex Lin',text);self.assertNotIn('schema',text);self.assertIsNone(w.tray);d.close()
    def test_bid_number_dialog_edits_stable_record(self):
        from test_finance import bid
        w=self.window;w.session.unlocked.add('bids');w.submit_bid(bid());row=w.session.pages['bids'][0]
        def accept():
            d=next(d for d in w.findChildren(QDialog) if d.isVisible());d.findChild(QSpinBox,'bid-number').setValue(10);d.findChild(QDialogButtonBox).button(QDialogButtonBox.StandardButton.Save).click()
        from PySide6.QtCore import QTimer
        QTimer.singleShot(0,accept);w.edit_bid(row);self.assertEqual(w.session.pages['bids'][0]['number'],10);self.assertEqual(w.session.pages['bids'][0]['id'],row['id']);w.tabs.setCurrentIndex(0);w.basic_tabs.setCurrentIndex(3);self.assertEqual(w.report_snapshot().rows[0][1],10)
    def test_removed_admin_fields_and_stage_setting_roundtrip(self):
        w=self.window;fields={k for k,v in w.payment_columns['admin_payments']};self.assertFalse(fields & {'estimated','paid','paid_tax','retention','submitted'});self.assertIn('tax',fields)
        w.session.pages['admin_payments']['allocations'][0].update(period=99,**{'0C12':'10','0C11':'20','0C14':'30'});w.session.pages['other_payments']=[transfer(period=3)];w.store.save(w.session.pages,0);loaded=w.store.load()['pages'];self.assertEqual(loaded['admin_payments']['allocation_revision'],2);self.assertEqual(effective_period(loaded,'first'),3);self.assertEqual(loaded['admin_payments']['allocations'][0]['period'],99)
    def test_export_snapshot_filters_and_ids_without_dirty(self):
        w=self.window;w.tabs.setCurrentIndex(4);w.diary_tabs.setCurrentIndex(0);w.session.pages['diary']=[dict(id='a',number=9,date='2026-10-20',description='甲',note=''),dict(id='b',number=1,date='2026-10-21',description='乙',note='')];w.populate_advanced('diary');w.session.saved=deepcopy(w.session.pages);w.diary_search.setText('乙');r=w.report_snapshot();self.assertEqual(r.rows,(('2026-10-21','乙',''),));self.assertEqual(len(w.report_snapshot(all_rows=True).rows),2);self.assertFalse(w.session.dirty())
    def test_payment_export_cutoff_and_immutable_snapshot(self):
        w=self.window;w.session.pages['other_payments']=[transfer(period=2),dict(transfer('second',10),amount='12.3456')];w.populate_billing();w.tabs.setCurrentIndex(2);w.billing_tabs.setCurrentIndex(1);w.payment_periods['other_payments'].setCurrentIndex(w.payment_periods['other_payments'].findData(2));r=w.report_snapshot();self.assertEqual(len(r.rows),1);w.session.pages['other_payments'][0]['amount']='888';self.assertIn('999',r.rows[0]);self.assertNotIn('888',r.rows[0]);self.assertEqual(len(w.report_snapshot(all_rows=True).rows),2)
    def test_calendar_export_has_all_month_days(self):
        w=self.window;w.tabs.setCurrentIndex(1);w.calendar_tabs.setCurrentWidget(w.calendar_page);w.calendar.setSelectedDate(QDate(2026,2,1));w.refresh_calendar();r=w.report_snapshot();self.assertTrue(r.calendar);self.assertEqual(len(r.rows),28);self.assertEqual(r.rows[0][0],'2026-02-01')
    def test_windows_tray_branch_option_and_minimize_fallback(self):
        from unittest.mock import MagicMock
        w=self.window;fake=MagicMock();fake.isVisible.return_value=True
        with patch('traaduoduo.desktop_features.sys.platform','win32'),patch('traaduoduo.desktop_features.QSystemTrayIcon') as tray:
            tray.isSystemTrayAvailable.return_value=True;tray.return_value=fake;w.install_tray();self.assertIs(w.tray,fake)
            with patch.object(w,'hide') as hidden:w.minimize_to_tray();self.assertTrue(hidden.called)
            with patch.object(w,'isMinimized',return_value=True),patch('traaduoduo.desktop_features.QTimer.singleShot') as later:w.changeEvent(QEvent(QEvent.Type.WindowStateChange));self.assertFalse(later.called)
            tray.isSystemTrayAvailable.return_value=False
            with patch.object(w,'isMinimized',return_value=True),patch('traaduoduo.desktop_features.QTimer.singleShot') as later:w.changeEvent(QEvent(QEvent.Type.WindowStateChange));self.assertFalse(later.called)
        w.tray=None
    def test_windows_atomic_lock_branch(self):
        from traaduoduo.instance import MacInstance
        root=Path(self.temp.name)
        with patch('traaduoduo.instance.sys.platform','win32'):
            first=MacInstance(root);second=MacInstance(root)
            try:self.assertTrue(first.owner);self.assertFalse(second.owner)
            finally:second.close();first.close()
            third=MacInstance(root);self.assertTrue(third.owner);third.close()
    def test_manual_save_waits_for_background_snapshot(self):
        w=self.window;w.toggle_lock('basic');w.fields['number'].setText('first');w.capture_basic();gate=Event();real=w.store.save
        def slow(*args):gate.wait(2);return real(*args)
        with patch.object(w.store,'save',side_effect=slow):
            w.start_background_save();w.fields['number'].setText('second');w.capture_basic();w.save();self.assertEqual(len(w._auto_after),1);gate.set();self.wait_io()
        self.assertEqual(w.store.load()['pages']['basic']['number'],'second');self.assertFalse(w.session.dirty())
    def test_table_editor_idle_does_not_force_commit(self):
        w=self.window;w.tabs.setCurrentIndex(4);w.diary_tabs.setCurrentIndex(0);w.session.unlocked.add('diary');w.commit_extra('diary',[dict(id='x',number=1,date='2026-10-22',description='內容',note='')]);w.update_actions();w.diary_table.setCurrentCell(0,1);w.diary_table.editItem(w.diary_table.item(0,1));support.APP.processEvents();self.assertTrue(w.input_activity.blocked());w.start_background_save();self.assertIsNone(w._auto_job);w.diary_search.setFocus();support.APP.processEvents()
    def test_allocation_dialog_has_no_manual_effective_period(self):
        from PySide6.QtCore import QTimer
        w=self.window;w.session.unlocked.add('admin_payments')
        def cancel():
            d=next(d for d in w.findChildren(QDialog) if d.isVisible());self.assertIsNone(d.findChild(QLineEdit,'allocation_period'));d.reject()
        QTimer.singleShot(0,cancel);w.edit_allocation(0)
    def test_real_schema9_readonly_import_preserves_cost_history(self):
        import subprocess,sys,os,hashlib
        from traaduoduo.importer import preview_database
        old=Path(__file__).resolve().parents[2]/'V0.7.0'
        if not old.is_dir():self.skipTest('V0.7.0 source required for actual-schema9 fixture')
        source=Path(self.temp.name)/'old.db';code="from traaduoduo.store import ProjectStore; import sys; p=ProjectStore.create(sys.argv[1],'舊庫'); p['pages']['costs']['budget'].update(award_total='100',unfinished_tax='5'); ProjectStore(sys.argv[1]).save(p['pages'],0)"
        subprocess.run([sys.executable,'-c',code,str(source)],cwd=old,env=dict(os.environ,PYTHONPATH=str(old),PYTHONDONTWRITEBYTECODE='1'),check=True);digest=hashlib.sha256(source.read_bytes()).hexdigest();preview=preview_database(source);self.assertFalse(preview['projects'][0]['errors']);result=self.window.store.import_preview(preview);loaded=ProjectStore(self.path,result['ids'][0]).load();self.assertEqual(loaded['metadata']['schema_version'],'10');self.assertEqual(loaded['pages']['costs']['budget']['historical_v070']['award_total'],'100');self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(),digest)
    def test_actual_export_dialog_pdf_xlsx_draft_and_cancel(self):
        from PySide6.QtCore import QTimer
        from PySide6.QtWidgets import QComboBox
        w=self.window;w.tabs.setCurrentIndex(4);w.diary_tabs.setCurrentIndex(0);w.session.unlocked.add('diary');w.commit_extra('diary',[dict(id='draft',number=1,date='2026-10-22',description='匯出草稿',note='')])
        for index,extension in [(0,'pdf'),(1,'xlsx')]:
            path=Path(self.temp.name)/('draft.'+extension)
            def accept():
                d=next(d for d in w.findChildren(QDialog) if d.isVisible());combos=d.findChildren(QComboBox);combos[0].setCurrentIndex(index);combos[-1].setCurrentIndex(1);d.findChild(QDialogButtonBox).button(QDialogButtonBox.StandardButton.Save).click()
            QTimer.singleShot(0,accept)
            with patch('traaduoduo.export_forms.QFileDialog.getSaveFileName',return_value=(str(path),'')):w.export_report_dialog()
            self.wait_io();self.assertTrue(path.is_file())
            if extension=='pdf':self.assertIn('未儲存草稿',PdfReader(path).pages[0].extract_text())
        self.assertEqual(w.store.load()['metadata']['revision'],'0');self.assertTrue(w.session.dirty())
        QTimer.singleShot(0,lambda:next(d for d in w.findChildren(QDialog) if d.isVisible()).reject())
        with patch('traaduoduo.export_forms.QFileDialog.getSaveFileName') as choose:w.export_report_dialog();self.assertFalse(choose.called)

class Export(unittest.TestCase):
    def setUp(self):
        import tempfile
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.report=Report('示範工程','期限與金額','截至2026-10-22',(Column('document_no','文號'),Column('due_date','日期','date'),Column('amount','金額','money'),Column('note','備註')), (('001-甲','2026-10-22','-1234.5678','=HYPERLINK("example")'),('002-乙','','0','中文長文測試')),draft=True)
    def tearDown(self):self.temp.cleanup()
    def test_xlsx_types_precision_and_formula_injection(self):
        path=self.root/'demo.xlsx';atomic_export(self.report,path,'xlsx');book=load_workbook(path);s=book.active;self.assertEqual(s.cell(7,1).value,'001-甲');self.assertEqual(s.cell(7,2).value.date(),date(2026,10,22));self.assertEqual(s.cell(7,3).value,-1234.5678);self.assertEqual(s.cell(7,3).number_format,'#,##0.00');self.assertEqual(s.cell(7,4).data_type,'s');self.assertFalse(any(c.data_type=='f' for row in s for c in row));book.close()
    def test_pdf_chinese_searchable_and_long_rows_repeat_headers(self):
        report=Report('示範工程','事記本','全部',(Column('date','日期','date'),Column('description','說明')),tuple(('2026-10-22','中文長文內容。'*70) for _ in range(30)))
        path=self.root/'long.pdf';atomic_export(report,path,'pdf');reader=PdfReader(path);self.assertGreater(len(reader.pages),1);self.assertIn('中文長文內容',reader.pages[0].extract_text());self.assertIn('日期',reader.pages[-1].extract_text());self.assertIn('說明',reader.pages[-1].extract_text())
    def test_failed_export_keeps_existing_target(self):
        path=self.root/'old.pdf';path.write_bytes(b'original')
        with self.assertRaises(ValueError):atomic_export(self.report,path,'pdf')
        with patch('traaduoduo.report_export.write_pdf',side_effect=OSError('full')):
            with self.assertRaises(OSError):atomic_export(self.report,path,'pdf',overwrite=True)
        self.assertEqual(path.read_bytes(),b'original');self.assertEqual(list(self.root.glob('.traa-export-*')),[])
    def test_wide_pdf_bands_keep_all_columns(self):
        report=Report('寬表','詳細表','全部',tuple(Column(str(i),'欄位'+str(i)) for i in range(16)),(tuple('內容'+str(i) for i in range(16)),));path=self.root/'wide.pdf';atomic_export(report,path,'pdf',landscape=True);text=''.join(p.extract_text() for p in PdfReader(path).pages)
        for i in range(16):self.assertIn('欄位'+str(i),text);self.assertIn('內容'+str(i),text)
