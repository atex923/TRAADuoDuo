"""Typed, immutable report snapshots with atomic local PDF/XLSX exports."""
from dataclasses import dataclass
from datetime import date,datetime
from decimal import Decimal
from pathlib import Path
from tempfile import NamedTemporaryFile
import os
from xml.sax.saxutils import escape
from . import APP_NAME,DISPLAY_VERSION

@dataclass(frozen=True)
class Column:
    key:str
    title:str
    kind:str='text'

@dataclass(frozen=True)
class Report:
    project:str
    title:str
    scope:str
    columns:tuple
    rows:tuple
    summary:tuple=()
    draft:bool=False
    calendar:bool=False

def display(value,kind):
    if value is None or value=='':return ''
    if kind in ('number','money') or isinstance(value,Decimal):
        try:return format(Decimal(str(value).replace(',','')),',.2f' if kind=='money' or isinstance(value,Decimal) else ',f')
        except ArithmeticError:return str(value)
    return str(value)

def atomic_export(report,path,kind,overwrite=False,paper='A4',landscape=False):
    target=Path(path).resolve()
    if target.exists() and not overwrite:raise ValueError('輸出檔已存在，須確認覆寫')
    if not report.columns:raise ValueError('請至少選擇一個匯出欄位')
    with NamedTemporaryFile(prefix='.traa-export-',suffix='.'+kind,dir=target.parent,delete=False) as file:temporary=Path(file.name)
    try:
        if kind=='xlsx':write_xlsx(report,temporary)
        elif kind=='pdf':write_pdf(report,temporary,paper,landscape)
        else:raise ValueError('未知匯出格式')
        if not temporary.stat().st_size:raise ValueError('輸出檔為空')
        if not overwrite and target.exists():raise ValueError('輸出目標已出現同名檔，請重新確認')
        os.replace(temporary,target)
    finally:
        if temporary.exists():temporary.unlink()
    return str(target)

def write_xlsx(report,path):
    # Runtime export belongs to the standalone Python application, using its
    # existing dependency; QA independently imports the result for inspection.
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    book=Workbook();sheet=book.active;sheet.title='報表';sheet.sheet_view.showGridLines=False
    def text_cell(row,col,value):
        cell=sheet.cell(row,col,str(value));cell.data_type='s';return cell
    text_cell(2,1,report.project+'／'+report.title);sheet.cell(2,1).font=Font(name='Arial',size=14,bold=True)
    text_cell(3,1,report.scope+('（未儲存草稿）' if report.draft else ''));text_cell(4,1,APP_NAME+' '+DISPLAY_VERSION)
    line=5
    for title,value in report.summary:
        text_cell(line,1,title)
        try:
            number=Decimal(str(value).replace(',','').replace('%',''));cell=sheet.cell(line,2,number/100 if str(value).endswith('%') else number);cell.number_format='0.00%' if str(value).endswith('%') else '#,##0.00'
        except ArithmeticError:text_cell(line,2,value)
        line+=1
    line+=1;header_row=line
    for col,column in enumerate(report.columns,1):
        cell=text_cell(line,col,column.title);cell.font=Font(name='Arial',size=11,bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='285C48');cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        width=48 if column.key in ('description','note','work') else 24 if column.kind=='text' else 20
        sheet.column_dimensions[get_column_letter(col)].width=width
    sheet.row_dimensions[line].height=32
    for values in report.rows:
        line+=1;longest=1
        for col,(column,value) in enumerate(zip(report.columns,values),1):
            cell=sheet.cell(line,col);cell.font=Font(name='Arial',size=11);cell.alignment=Alignment(horizontal='right' if column.kind in ('number','money') else 'left',vertical='top',wrap_text=True,indent=1)
            if value=='' or value is None:cell.value=None
            elif column.kind in ('number','money') or isinstance(value,(Decimal,int,float)):
                try:number=Decimal(str(value).replace(',',''))
                except ArithmeticError:cell.value=str(value);cell.data_type='s';continue
                if not number.is_finite():raise ValueError('金額不可為非有限值')
                cell.value=number;cell.number_format='#,##0.00' if column.kind=='money' else '#,##0' if number==number.to_integral_value() else '#,##0.########'
            elif isinstance(value,date):cell.value=value;cell.number_format='yyyy-mm-dd hh:mm' if isinstance(value,datetime) else 'yyyy-mm-dd'
            elif column.kind=='date' or column.kind=='issued' and len(str(value))==10:cell.value=date.fromisoformat(str(value));cell.number_format='yyyy-mm-dd'
            elif column.kind in ('datetime','issued'):cell.value=datetime.strptime(str(value),'%Y-%m-%d %H:%M');cell.number_format='yyyy-mm-dd hh:mm'
            else:cell.value=str(value);cell.data_type='s'
            visual_length=sum(2 if ord(ch)>255 else 1 for ch in str(value));width=max(12,int(sheet.column_dimensions[get_column_letter(col)].width)-2)
            longest=max(longest,(visual_length+width-1)//width+str(value).count('\n'))
        sheet.row_dimensions[line].height=max(22,min(400,16*longest))
    sheet.freeze_panes=f'A{header_row+1}';sheet.auto_filter.ref=f'A{header_row}:{get_column_letter(len(report.columns))}{max(header_row,line)}'
    sheet.print_title_rows=f'{header_row}:{header_row}';sheet.page_setup.orientation='landscape' if len(report.columns)>5 else 'portrait';sheet.sheet_properties.pageSetUpPr.fitToPage=True;sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
    book.save(path);book.close()

def pdf_font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    font='TRAA-CJK'
    if font in pdfmetrics.getRegisteredFontNames():return font
    for path in ['/System/Library/Fonts/Supplemental/Arial Unicode.ttf','C:/Windows/Fonts/msjh.ttc','C:/Windows/Fonts/mingliu.ttc','/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc']:
        if Path(path).is_file():
            try:pdfmetrics.registerFont(TTFont(font,path));return font
            except Exception:pass
    pdfmetrics.registerFont(UnicodeCIDFont('MSung-Light'));return 'MSung-Light'

def write_pdf(report,path,paper='A4',wide=False):
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,LongTable,TableStyle,PageBreak
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4,A3,landscape
    from reportlab.lib import colors
    size=A3 if paper=='A3' else A4
    if wide:size=landscape(size)
    font=pdf_font();body=ParagraphStyle('body',fontName=font,fontSize=10,leading=14,wordWrap='CJK');heading=ParagraphStyle('heading',parent=body,fontSize=14,leading=20)
    def p(value):return Paragraph(escape(str(value)).replace('\n','<br/>'),body)
    doc=SimpleDocTemplate(str(path),pagesize=size,rightMargin=30,leftMargin=30,topMargin=34,bottomMargin=36,title=report.title,author=APP_NAME)
    content=[Paragraph(escape(report.project+'／'+report.title),heading),p(report.scope+('（未儲存草稿）' if report.draft else '')),Spacer(1,8)]
    for label,value in report.summary:content.append(p(label+'：'+str(value)))
    if report.summary:content.append(Spacer(1,8))
    # Wide tables become readable column bands, retaining the first identifying
    # column, rather than shrinking all sixteen columns into tiny print.
    maximum=7 if wide or paper=='A3' else 4
    groups=[list(range(len(report.columns)))] if len(report.columns)<=maximum else [[0]+list(range(i,min(i+maximum-1,len(report.columns)))) for i in range(1,len(report.columns),maximum-1)]
    if report.calendar:
        import calendar
        content.extend([p('各格先列日期，再列事件與累計工作日數。'),Spacer(1,6)])
        first=date.fromisoformat(report.rows[0][0]);by_date={r[0]:r for r in report.rows}
        matrix=[[p('星期'+d) for d in '一二三四五六日']]
        for week in calendar.monthcalendar(first.year,first.month):
            cells=[]
            for day in week:
                r=by_date.get(date(first.year,first.month,day).isoformat()) if day else None
                cells.append(p('' if not r else '\n'.join([str(day),*[str(v) for v in r[2:] if v!='']])))
            matrix.append(cells)
        table=LongTable(matrix,colWidths=[doc.width/7]*7,repeatRows=1,splitInRow=1);table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#DDEBDD')),('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),.3,colors.HexColor('#BBCABB')),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,1),(-1,-1),18)]));content.append(table)
    for band,indices in enumerate([] if report.calendar else groups):
        if band:content.extend([PageBreak(),p(report.project+'／'+report.title+f'／欄組 {band+1}')])
        matrix=[[p(report.columns[i].title) for i in indices]]
        matrix.extend([[p(display(values[i],report.columns[i].kind)) for i in indices] for values in report.rows])
        weights=[3 if report.columns[i].key in ('work','description','note') else 1.5 if report.columns[i].kind=='text' else 1 for i in indices];unit=doc.width/sum(weights)
        table=LongTable(matrix,colWidths=[unit*w for w in weights],repeatRows=1,splitByRow=1,splitInRow=1)
        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#DDEBDD')),('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,0),.6,colors.HexColor('#52765A')),('LINEBELOW',(0,1),(-1,-1),.2,colors.HexColor('#CFD7CF')),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]));content.append(table)
    def footer(canvas,doc):
        canvas.setFont(font,8);canvas.drawString(30,19,APP_NAME+' '+DISPLAY_VERSION);canvas.drawRightString(size[0]-30,19,str(doc.page))
    doc.build(content,onFirstPage=footer,onLaterPages=footer)
